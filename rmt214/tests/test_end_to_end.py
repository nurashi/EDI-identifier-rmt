"""Integration test: two tiny hand-made PDFs in the usual guide layout are run
through extraction -> alignment -> candidate overview -> blind coder sheets.
The PDFs are TEST FIXTURES written by us, not partner guides and not research data."""

import sys
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract  # noqa: E402
from align import units_for_pairs  # noqa: E402
from annotate import export  # noqa: E402
from prelabel import prelabel  # noqa: E402

FIXTURE_A = [
    "B10      Beginning Segment for Transportation Carrier Shipment Status Message",
    "Pos: 020      Max: 1",
    "Heading - Mandatory",
    "User Option (Usage): Must use",
    "Ref    Id    Element Name                    Req   Type   Min/Max   Usage",
    "B1001  127   Reference Identification          M     AN     1/30      Must use",
    "B1002  145   Shipment Identification Number    M     AN     1/30      Must use",
    "B1003  140   Standard Carrier Alpha Code       M     ID     2/4       Must use",
]
FIXTURE_B = [
    "B10      Beginning Segment for Transportation Carrier Shipment Status Message",
    "Pos: 020      Max: 1",
    "Heading - Mandatory",
    "User Option (Usage): Must use",
    "Ref    Id    Element Name                    Req   Type   Min/Max   Usage",
    "B1001  127   Reference Identification          M     AN     1/30      Must use",
    "B1002  145   Shipment Identification Number    M     AN     1/30      Used",
    "B1003  140   Standard Carrier Alpha Code       M     ID     2/4       Must use",
]


def write_pdf(path, lines):
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Courier", 8)
    y = 740
    for ln in lines:
        c.drawString(30, y, ln)
        y -= 12
    c.save()


def test_pdf_to_coder_sheet(tmp_path):
    pa, pb = tmp_path / "a.pdf", tmp_path / "b.pdf"
    write_pdf(pa, FIXTURE_A); write_pdf(pb, FIXTURE_B)
    import pandas as pd
    frames = []
    for g, p in (("TA", pa), ("TB", pb)):
        df, unparsed = extract.parse_guide(g, extract.pdf_pages(p))
        assert not df.empty
        frames.append(df)
    c = pd.concat(frames, ignore_index=True)
    assert set(c["page"]) == {1}                       # page evidence kept
    u = prelabel(units_for_pairs(c)).set_index("x12_position")
    assert u.loc["H020 B10 / B1002", "cand_status"] == "DIFFERENT"   # Must use vs Used
    assert u.loc["H020 B10 / B1001", "cand_status"] == "NONE"
    blind = u.reset_index().drop(columns=[x for x in u.columns if x.startswith("cand_")])
    out = tmp_path / "coder1.xlsx"
    export(blind, out, "coder1", {"TA": "fixture A", "TB": "fixture B"})
    ws = load_workbook(out)["units"]
    headers = [h.value for h in ws[1]]
    assert "status" in headers and not any(str(h).startswith("cand_") for h in headers)
