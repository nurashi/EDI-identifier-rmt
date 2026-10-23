"""Export blind coder sheets (Objective 3).

Each coder gets an identical workbook: same units, same random order, the
evidence of both guides, empty label columns with drop-down validation, and the
codebook. Rule-based candidate labels are deliberately excluded.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config
from align import attr_cols

CODEBOOK_ROWS = [
    ("status", "NONE", "Both guides document this position and every stated attribute is equivalent; "
     "differences are only wording, layout or element names."),
    ("status", "DIFFERENT", "Both guides document the position and at least one stated requirement differs "
     "in a way a sender would have to handle differently."),
    ("status", "UNKNOWN", "The position or a deciding attribute is not documented in one guide, or the text "
     "is too ambiguous to decide. Absence is never read as 'Not used'."),
    ("category", "USAGE", "Segment/element usage or requirement differs (e.g. Must use vs Used vs Not used)."),
    ("category", "CODES", "The sets of permitted code values differ (added, removed or restricted codes)."),
    ("category", "REPETITION", "Max use of a segment or repeat of a loop differs."),
    ("category", "LENGTH_FORMAT", "Min/max length or data type differs."),
    ("category", "CONDITION", "A conditional or semantic instruction (if/when/only) applies in one guide "
     "and is absent or different in the other."),
    ("rule", "scope", "If a requirement is stated only for a non-truckload mode, mark UNKNOWN and write 'scope' in note."),
    ("rule", "evidence", "Base every label on the shown page text; if you open the PDF, cite the page in note."),
    ("rule", "independence", "Do not discuss units with the other coder until both sheets are submitted."),
]

FILL_HEAD = PatternFill("solid", fgColor="DDE6F0")
FILL_CODE = PatternFill("solid", fgColor="FFF4CC")


def export(units: pd.DataFrame, path, coder: str, guide_labels: dict):
    order = units.sample(frac=1.0, random_state=config.RANDOM_SEED)   # same order for both coders
    ev = ["unit_id", "guide_pair", "x12_position", "documented_in"] + attr_cols(order, "A") + attr_cols(order, "B")
    label_cols = ["status"] + [f"cat_{c}" for c in config.CATEGORIES] + ["note"]

    wb = Workbook()
    ws = wb.active
    ws.title = "units"
    headers = ev + label_cols
    ws.append(headers)
    for _, r in order.iterrows():
        ws.append([("" if pd.isna(r.get(c)) else r.get(c)) for c in ev] + [""] * len(label_cols))

    n = len(order) + 1
    dv_s = DataValidation(type="list", formula1='"' + ",".join(config.STATUS) + '"', allow_blank=True)
    dv_c = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv_s); ws.add_data_validation(dv_c)
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        cell = ws.cell(1, i)
        cell.font = Font(bold=True)
        cell.fill = FILL_CODE if h in label_cols else FILL_HEAD
        ws.column_dimensions[col].width = 38 if h.endswith("__ev__A") or h.endswith("__ev__B") else 14
        if h == "status":
            dv_s.add(f"{col}2:{col}{n}")
        elif h.startswith("cat_"):
            dv_c.add(f"{col}2:{col}{n}")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"

    cb = wb.create_sheet("codebook")
    cb.append(["type", "label", "operational definition"])
    for r in CODEBOOK_ROWS:
        cb.append(list(r))
    cb.column_dimensions["C"].width = 110
    info = wb.create_sheet("about")
    info.append(["coder", coder])
    info.append(["units", len(order)])
    info.append(["transaction/version", f"{config.TRANSACTION} / {config.VERSION}"])
    for k, v in guide_labels.items():
        info.append([f"guide {k}", v])
    info.append(["note", "Columns ending __A belong to the first guide of guide_pair, __B to the second."])
    wb.save(path)
